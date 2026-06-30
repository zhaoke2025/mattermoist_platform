from __future__ import annotations

import csv
import io
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


MAX_EXTRACTED_CHARS = 120_000


class ExtractionError(Exception):
    pass


def clamp_text(text: str, limit: int = MAX_EXTRACTED_CHARS) -> str:
    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = "\n".join(line for line in text.splitlines() if line.strip())
    if len(text) <= limit:
        return text
    return text[:limit] + "\n\n[Content truncated because it is too long]"


def extract_text(file_path: Path, original_name: str) -> str:
    suffix = Path(original_name).suffix.lower()
    if suffix == ".pdf":
        return clamp_text(_extract_pdf(file_path))
    if suffix == ".docx":
        return clamp_text(_extract_docx(file_path))
    if suffix == ".xlsx":
        return clamp_text(_extract_xlsx(file_path))
    if suffix == ".pptx":
        return clamp_text(_extract_pptx(file_path))
    if suffix in {".txt", ".md", ".csv", ".log"}:
        return clamp_text(_extract_text_like(file_path, suffix))
    raise ExtractionError(f"Unsupported file type: {suffix or 'unknown'}")


def _extract_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    chunks: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            chunks.append(f"--- Page {index} ---\n{text}")
    if not chunks:
        raise ExtractionError("No readable text was extracted from the PDF. It may be scanned images.")
    return "\n\n".join(chunks)


def _extract_docx(file_path: Path) -> str:
    doc = Document(str(file_path))
    chunks: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text)
    for table_index, table in enumerate(doc.tables, start=1):
        chunks.append(f"--- Table {table_index} ---")
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    if not chunks:
        raise ExtractionError("No readable text was extracted from the DOCX file.")
    return "\n".join(chunks)


def _extract_xlsx(file_path: Path) -> str:
    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets:
        chunks.append(f"--- Sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                chunks.append(" | ".join(values))
    if not chunks:
        raise ExtractionError("No readable text was extracted from the XLSX file.")
    return "\n".join(chunks)


def _extract_pptx(file_path: Path) -> str:
    prs = Presentation(str(file_path))
    chunks: list[str] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        slide_text: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        if slide_text:
            chunks.append(f"--- Slide {slide_index} ---\n" + "\n".join(slide_text))
    if not chunks:
        raise ExtractionError("No readable text was extracted from the PPTX file.")
    return "\n\n".join(chunks)


def _extract_text_like(file_path: Path, suffix: str) -> str:
    raw = file_path.read_bytes()
    text = raw.decode("utf-8", errors="replace")
    if suffix == ".csv":
        return _format_csv(text)
    if not text.strip():
        raise ExtractionError("No readable text was extracted from the text file.")
    return text


def _format_csv(text: str) -> str:
    reader = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in reader)
